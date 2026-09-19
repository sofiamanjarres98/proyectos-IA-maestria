from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from src.core.holidays_colombia import colombia_holidays
from src.optimization.scheduler import base_demand, month_dates


SHIFT_FILLS = {
    "M": "DBEAFE",
    "T": "FDE68A",
    "N": "DDD6FE",
    "C": "BBF7D0",
    "L": "F3F4F6",
}


def build_template(path: Path, year: int, month: int) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "PERSONAL"
    ws.append(
        [
            "internal_code",
            "full_name",
            "staff_type",
            "monthly_salary",
            "ops_day_rate",
            "ops_night_rate",
            "target_monthly_hours",
            "max_monthly_hours",
            "desired_monthly_hours",
        ]
    )
    ws.append(["D001", "Medico Directo 1", "DIRECTO", 5970000, "", "", 210, 240, ""])
    ws.append(["OPS1", "Medico OPS 1", "OPS", "", 30000, 40000, "", 120, 96])

    ws = wb.create_sheet("DISPONIBILIDAD")
    ws.append(["internal_code", "date", "shift", "available"])
    for day in month_dates(year, month):
        ws.append(["OPS1", day.isoformat(), "N", True])

    ws = wb.create_sheet("DEMANDA")
    ws.append(["date", "shift", "required"])
    for row in base_demand(year, month).itertuples(index=False):
        ws.append([row.date.isoformat(), row.shift, row.required])

    ws = wb.create_sheet("NOVEDADES")
    ws.append(["internal_code", "start_date", "end_date", "event_type", "blocks_assignment"])
    ws.append(["D001", f"{year}-{month:02d}-03", f"{year}-{month:02d}-05", "VACACIONES", True])

    ws = wb.create_sheet("PARAMETROS")
    ws.append(["name", "value", "unit", "requires_admin_validation"])
    ws.append(["weekly_max_hours", 42, "hours", True])
    ws.append(["monthly_payroll_hour_divisor", 210, "hours", True])
    ws.append(["max_shift_hours", 12, "hours", True])

    ws = wb.create_sheet("FESTIVOS")
    ws.append(["date", "name", "source", "admin_verified"])
    for holiday in colombia_holidays(year):
        ws.append([holiday.date.isoformat(), holiday.name, holiday.source, False])

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E5E7EB")
        sheet.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_schedule(path: Path, schedule: pd.DataFrame, summary: pd.DataFrame, alerts: pd.DataFrame) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pivot = schedule.pivot_table(index=["tipo", "nombre"], columns="dia", values="shift", aggfunc="first").reset_index()
        pivot.to_excel(writer, sheet_name="CALENDARIO", index=False)
        summary.to_excel(writer, sheet_name="RESUMEN", index=False)
        alerts.to_excel(writer, sheet_name="VALIDACIONES", index=False)
    wb = Workbook()
    # Reopen through openpyxl via pandas output.
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["CALENDARIO"]
    for row in ws.iter_rows(min_row=2, min_col=3):
        for cell in row:
            if cell.value in SHIFT_FILLS:
                cell.fill = PatternFill("solid", fgColor=SHIFT_FILLS[cell.value])
    wb.save(path)
    return path
